import pptx
import msoffcrypto
import io
import os
import re
import json
from PIL import Image

BASE_DIR = "/home/longp/projects/tech-ws-app"
DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_IMG_DIR = os.path.join(DATA_DIR, "images/exams")
os.makedirs(OUTPUT_IMG_DIR, exist_ok=True)
os.makedirs(os.path.join(BASE_DIR, "images/exams"), exist_ok=True)

circle_map = {'①': 'A', '②': 'B', '③': 'C', '④': 'D', '⑤': 'E', '⑥': 'F'}
digit_to_alpha = {'1': 'A', '2': 'B', '3': 'C', '4': 'D', '5': 'E', '6': 'F'}

def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    text = re.sub(r'\r\n|\r', '\n', text)
    lines = [re.sub(r'[ \t]+', ' ', l).strip() for l in text.split('\n')]
    return '\n'.join(lines).strip()

def normalize_q(text):
    if not text: return ""
    t = text.lower()
    t = re.sub(r'^\d+[\.\s\:\)]+', '', t) # remove leading question number
    t = re.sub(r'[^\w\s]', '', t)
    t = re.sub(r'\s+', ' ', t).strip()
    return t

def parse_presentation(fpath, pw=None):
    if pw:
        with open(fpath, "rb") as fh:
            of = msoffcrypto.OfficeFile(fh)
            of.load_key(password=pw)
            s = io.BytesIO()
            of.decrypt(s)
            s.seek(0)
            return pptx.Presentation(s)
    else:
        return pptx.Presentation(fpath)

# Extract images recursively
def get_slide_images(slide):
    imgs = []
    def recurse(sh, p_left=0, p_top=0):
        if sh.shape_type == 13: # PICTURE
            # Check size to skip tiny noise
            if sh.width > 20000 and sh.height > 20000:
                imgs.append({
                    "shape": sh,
                    "left": sh.left + p_left,
                    "top": sh.top + p_top,
                    "width": sh.width,
                    "height": sh.height,
                    "blob": sh.image.blob,
                    "ext": sh.image.ext.lower()
                })
        elif sh.shape_type == 6: # GROUP
            for sub in sh.shapes:
                recurse(sub, p_left + sh.left, p_top + sh.top)
                
    for sh in slide.shapes:
        recurse(sh)
    return imgs

print("Module extract_all_exams_complete loaded successfully.")
