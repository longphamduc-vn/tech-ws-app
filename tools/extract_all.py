import openpyxl
import pptx
import os
import io
import json
import re
from PIL import Image

BASE_DIR = "/home/longp/projects/tech-ws-app"
DATA_DIR = os.path.join(BASE_DIR, "data")
IMAGES_DIR = os.path.join(DATA_DIR, "images")
SUBJECTS_DIR = os.path.join(DATA_DIR, "questions_by_subject")
EXAMS_DIR = os.path.join(DATA_DIR, "exams")

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(IMAGES_DIR, exist_ok=True)
os.makedirs(SUBJECTS_DIR, exist_ok=True)
os.makedirs(EXAMS_DIR, exist_ok=True)

# Helper clean text
def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    text = re.sub(r'\r\n|\r', '\n', text)
    # clean extra whitespace within lines
    lines = [re.sub(r'[ \t]+', ' ', l).strip() for l in text.split('\n')]
    return '\n'.join(lines).strip()

# ==============================================================================
# 1. TRÍCH XUẤT MASTER BANK TỪ EXCEL (398 CÂU)
# ==============================================================================
wb_path = os.path.join(BASE_DIR, "WS/05. Quan ly dao tao/Bang diem dao tao va on thi White Star.xlsx")
wb = openpyxl.load_workbook(wb_path)

category_meta = {
    'Sửa-Điện-Điện tử Basic': {'code': 'electric', 'name': 'Điện - Điện tử cơ bản', 'prefix': 'WS_ELE'},
    'Sửa-PLC Basic': {'code': 'plc', 'name': 'PLC cơ bản', 'prefix': 'WS_PLC'},
    'Sửa-Máy Basic': {'code': 'machine', 'name': 'Linh kiện máy cơ bản', 'prefix': 'WS_MCH'},
    'Sửa-Khí nén Basic': {'code': 'pneumatics', 'name': 'Khí nén cơ bản', 'prefix': 'WS_PNE'}
}

excel_row_images = {}
for sheet_name, meta in category_meta.items():
    sub_code = meta['code']
    cat_img_dir = os.path.join(IMAGES_DIR, sub_code)
    os.makedirs(cat_img_dir, exist_ok=True)
    sheet = wb[sheet_name]
    images = getattr(sheet, '_images', [])
    
    row_to_q = {}
    for r in range(4, sheet.max_row + 1):
        q_num = sheet.cell(r, 4).value
        if q_num is not None:
            try:
                row_to_q[r] = int(q_num)
            except Exception:
                row_to_q[r] = r
                
    img_by_row = {}
    for img in images:
        anchor = getattr(img, 'anchor', None)
        if hasattr(anchor, '_from'):
            r = anchor._from.row + 1
            c = anchor._from.col + 1
            img_by_row.setdefault(r, []).append((c, img))
            
    for r, img_list in img_by_row.items():
        q_num = row_to_q.get(r, r)
        img_list.sort(key=lambda x: x[0])
        rel_paths = []
        for idx, (c, img) in enumerate(img_list):
            ext = img.format.lower()
            if ext == 'jpeg': ext = 'jpg'
            suffix = f"_{idx+1}" if len(img_list) > 1 else ""
            filename = f"q_{q_num:03d}{suffix}.{ext}"
            file_path = os.path.join(cat_img_dir, filename)
            data = img._data()
            with open(file_path, "wb") as f:
                f.write(data)
            rel_paths.append(f"images/{sub_code}/{filename}")
        excel_row_images[(sub_code, r)] = rel_paths

missing_electric_questions = {
    3: "Dựa vào hình ảnh về sự kích thích và phóng xạ của nguyên tử, đâu không phải là hiện tượng được thể hiện?",
    4: "Quan sát sơ đồ mức năng lượng, mức năng lượng cơ bản (nền) thấp nhất của electron được gọi là gì?",
    14: "Dựa vào hình ảnh, khi electron hấp thụ năng lượng và chuyển dịch lên mức quỹ đạo cao hơn, hiện tượng này được gọi là gì?",
    17: "Dạng xung nhiễu được biểu diễn trong hình ảnh dưới đây tương ứng với loại nhiễu nào?",
    23: "Cho đồ thị sóng trên máy dao động ký (Oscilloscope) với biên độ cực đại Vp = 3V. Giá trị hiệu dụng (Vrms) của điện áp là bao nhiêu? (Vrms = 0.707 × Vp)",
    24: "Cho đồ thị sóng trên máy dao động ký với biên độ cực đại Vp = 3V. Giá trị trung bình (Vavg) của điện áp là bao nhiêu? (Vavg = 0.637 × Vp)",
    25: "Trên máy dao động ký với cài đặt trục thời gian M = 500μs/ô, một chu kỳ sóng chiếm 4 ô. Chu kỳ T của sóng là bao nhiêu?",
    26: "Trên máy dao động ký với thang đo 1.00V/ô, độ lệch đỉnh - đỉnh của sóng là 6 ô. Giá trị điện áp đỉnh - đỉnh (Vp-p) là bao nhiêu?",
    32: "Chu kỳ của điện áp xoay chiều là T = 2[ms]. Tần số (f) của dòng điện là bao nhiêu? (f = 1/T)",
    33: "Tần số sóng vô tuyến là f = 500[Hz], vận tốc truyền sóng v = 3×10⁸ m/s. Bước sóng λ là bao nhiêu? (λ = v / f)"
}

master_questions = []
by_subject = {'electric': [], 'plc': [], 'machine': [], 'pneumatics': []}

for sheet_name, meta in category_meta.items():
    sheet = wb[sheet_name]
    sub_code = meta['code']
    sub_name = meta['name']
    prefix = meta['prefix']
    
    for r in range(4, sheet.max_row + 1):
        q_num = sheet.cell(r, 4).value
        q_text = clean_text(sheet.cell(r, 5).value)
        ans = clean_text(sheet.cell(r, 10).value).upper()
        module = clean_text(sheet.cell(r, 2).value)
        difficulty = clean_text(sheet.cell(r, 3).value)
        author = clean_text(sheet.cell(r, 11).value)
        note = clean_text(sheet.cell(r, 12).value)
        
        if not q_text and not ans and not q_num:
            continue
        try:
            q_num_int = int(q_num)
        except Exception:
            continue
            
        if not q_text and sub_code == 'electric' and q_num_int in missing_electric_questions:
            q_text = missing_electric_questions[q_num_int]
            
        if not q_text:
            continue
            
        opts = []
        col_letters = ['A', 'B', 'C', 'D']
        for idx, col_idx in enumerate([6, 7, 8, 9]):
            opt_val = clean_text(sheet.cell(r, col_idx).value)
            if opt_val:
                opts.append({
                    "key": col_letters[idx],
                    "text": opt_val
                })
                
        if len(opts) == 0:
            opts = [
                {"key": "A", "text": "Đáp án A / Hình A (Xem sơ đồ minh họa)"},
                {"key": "B", "text": "Đáp án B / Hình B (Xem sơ đồ minh họa)"},
                {"key": "C", "text": "Đáp án C / Hình C (Xem sơ đồ minh họa)"},
                {"key": "D", "text": "Đáp án D / Hình D (Xem sơ đồ minh họa)"}
            ]
                
        images = excel_row_images.get((sub_code, r), [])
        valid_ans = ""
        for c in ['A', 'B', 'C', 'D']:
            if c in ans:
                valid_ans = c
                break
                
        q_id = f"{prefix}_{q_num_int:03d}"
        q_obj = {
            "id": q_id,
            "source": "Bang diem dao tao va on thi White Star.xlsx",
            "subject": sub_name,
            "subject_code": sub_code,
            "module": module,
            "difficulty": difficulty if difficulty in ["Dễ", "Trung bình", "Khó"] else "Trung bình",
            "question_number": q_num_int,
            "question": q_text,
            "options": opts,
            "correct_answer": valid_ans if valid_ans else ans,
            "images": images,
            "author": author,
            "note": note
        }
        master_questions.append(q_obj)
        by_subject[sub_code].append(q_obj)

# ==============================================================================
# 2. TRÍCH XUẤT REVIEW BANK TỪ PPTX (486 CÂU CÓ ĐÁP ÁN ĐỎ)
# ==============================================================================
review_prs_path = os.path.join(BASE_DIR, "WS/03. De on tap & Ngan hang cau hoi/Ngan hang cau hoi on tap White Star - Dap an.pptx")
review_prs = pptx.Presentation(review_prs_path)
review_img_dir = os.path.join(IMAGES_DIR, "review_bank")
os.makedirs(review_img_dir, exist_ok=True)

review_questions = []

circle_map = {'①': 'A', '②': 'B', '③': 'C', '④': 'D'}

for s_idx, slide in enumerate(review_prs.slides):
    slide_num = s_idx + 1
    subj = "Chung"
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text.strip()
            if t.startswith("[") and t.endswith("]"):
                subj = t.strip("[] ")
                
    # Save slide images if any
    slide_images = []
    pic_idx = 0
    for sh in slide.shapes:
        if sh.shape_type == 13:
            pic_idx += 1
            img_fname = f"slide_{slide_num:02d}_img_{pic_idx}.png"
            img_path = os.path.join(review_img_dir, img_fname)
            try:
                with open(img_path, "wb") as f:
                    f.write(sh.image.blob)
                slide_images.append(f"images/review_bank/{img_fname}")
            except Exception:
                pass
                
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text.strip()
        if not t: continue
        
        m = re.match(r'^(?:※[^\n]+\n)?(\d+)[\.\s]+(.+)', t, re.DOTALL)
        if m:
            q_num = int(m.group(1))
            body = m.group(2).strip()
            
            # Find red text
            red_texts = []
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.color and r.font.color.type == 1:
                        rgb = str(r.font.color.rgb).upper()
                        if rgb in ["FF0000", "C00000", "ED1C24", "E30613"]:
                            red_texts.append(r.text.strip())
            red_ans_text = " ".join(red_texts).strip()
            
            # Parse question and options
            opts = []
            q_title = ""
            
            if any(c in body for c in ['①', '②', '③', '④']):
                parts = re.split(r'([①②③④])', body)
                q_title = parts[0].strip()
                for i in range(1, len(parts), 2):
                    ckey = parts[i]
                    cval = parts[i+1].strip() if i+1 < len(parts) else ""
                    opts.append({
                        "key": circle_map.get(ckey, ckey),
                        "symbol": ckey,
                        "text": clean_text(cval)
                    })
            else:
                lines = [l.strip() for l in body.split('\n') if l.strip()]
                if lines:
                    q_title = lines[0]
                    opt_lines = lines[1:]
                    for idx, l in enumerate(opt_lines):
                        m_opt = re.match(r'^([A-D1-4])[\.\s]+(.*)', l)
                        if m_opt:
                            k = m_opt.group(1)
                            # map 1->A, 2->B...
                            if k in ['1','2','3','4']:
                                k = chr(ord('A') + int(k) - 1)
                            opts.append({"key": k, "text": clean_text(m_opt.group(2))})
                        else:
                            k = chr(ord('A') + idx) if idx < 4 else f"Opt_{idx+1}"
                            opts.append({"key": k, "text": clean_text(l)})
                            
            # Determine correct answer key based on red_ans_text
            corr_key = ""
            if red_ans_text:
                for opt in opts:
                    if red_ans_text.lower() in opt['text'].lower() or opt['text'].lower() in red_ans_text.lower():
                        corr_key = opt['key']
                        break
                if not corr_key:
                    # check if red text contains key directly like "B. Ion hóa"
                    m_key = re.match(r'^([A-D①-④1-4])', red_ans_text)
                    if m_key:
                        k_char = m_key.group(1)
                        corr_key = circle_map.get(k_char, k_char)
                        if corr_key in ['1','2','3','4']:
                            corr_key = chr(ord('A') + int(corr_key) - 1)
                            
            review_questions.append({
                "id": f"WS_REV_S{slide_num:02d}_Q{q_num:02d}",
                "source": "Ngan hang cau hoi on tap White Star - Dap an.pptx",
                "slide": slide_num,
                "subject": subj,
                "question_number": q_num,
                "question": clean_text(q_title if q_title else t.splitlines()[0]),
                "options": opts,
                "correct_answer": corr_key,
                "answer_text": red_ans_text,
                "images": slide_images
            })

# ==============================================================================
# 3. TRÍCH XUẤT 5 BỘ ĐỀ THI THỬ CHUẨN (MOCK EXAMS 01 - 04, ĐỀ 01: 80 CÂU MỖI ĐỀ)
# ==============================================================================
mock_exams = []

# Mock Exam De 01
de01_answers = {
    1: 'D', 2: 'D', 3: 'C', 4: 'A', 5: 'B', 6: 'B', 7: 'B', 8: 'C', 9: 'A', 10: 'C',
    11: 'D', 12: 'A', 13: 'C', 14: 'B', 15: 'D', 16: 'A', 17: 'D', 18: 'B', 19: 'C', 20: 'A',
    21: 'A', 22: 'A', 23: 'C', 24: 'A', 25: 'B', 26: 'D', 27: 'B', 28: 'C', 29: 'A', 30: 'C',
    31: 'D', 32: 'D', 33: 'C', 34: 'B', 35: 'A', 36: 'C', 37: 'B', 38: 'C', 39: 'D', 40: 'A',
    41: 'D', 42: 'C', 43: 'B', 44: 'A', 45: 'A', 46: 'D', 47: 'C', 48: 'A', 49: 'C', 50: 'A',
    51: 'A', 52: 'B', 53: 'C', 54: 'D', 55: 'B', 56: 'A', 57: 'C', 58: 'B', 59: 'A', 60: 'C',
    61: 'A', 62: 'A', 63: 'C', 64: 'C', 65: 'D', 66: 'A', 67: 'D', 68: 'D', 69: 'C', 70: 'A',
    71: 'D', 72: 'B', 73: 'C', 74: 'B', 75: 'A', 76: 'C', 77: 'C', 78: 'B', 79: 'B', 80: 'D'
}

def parse_exam_pptx(pptx_path, exam_id, exam_title, manual_answers=None):
    prs = pptx.Presentation(pptx_path)
    exam_questions = []
    
    # Answers map
    answers_map = {}
    if manual_answers:
        answers_map = manual_answers.copy()
    else:
        # Check last slide for tables
        last_s = prs.slides[-1]
        for sh in last_s.shapes:
            if sh.has_table:
                # determine section or columns
                rows = list(sh.table.rows)
                header = [c.text.strip() for c in rows[0].cells]
                for r in rows[1:]:
                    vals = [c.text.strip() for c in r.cells]
                    if len(vals) >= 2 and vals[0].isdigit():
                        q_i = int(vals[0])
                        a_i = vals[1]
                        # map 1->A, 2->B...
                        if a_i in ['1','2','3','4']:
                            a_i = chr(ord('A') + int(a_i) - 1)
                        answers_map[q_i] = a_i

    # Parse slides 0 to 7 (8 question slides)
    for s_idx, slide in enumerate(prs.slides):
        if s_idx >= 8:
            break
        sec_name = ""
        for sh in slide.shapes:
            if sh.has_text_frame:
                t = sh.text.strip()
                if t.startswith("[") and t.endswith("]"):
                    sec_name = t.strip("[] ")
                    
        slide_imgs = []
        for sh in slide.shapes:
            if sh.shape_type == 13:
                img_name = f"{exam_id}_s{s_idx+1}_{len(slide_imgs)+1}.png"
                img_p = os.path.join(IMAGES_DIR, "exams")
                os.makedirs(img_p, exist_ok=True)
                with open(os.path.join(img_p, img_name), "wb") as f:
                    f.write(sh.image.blob)
                slide_imgs.append(f"images/exams/{img_name}")

        for sh in slide.shapes:
            if not sh.has_text_frame: continue
            text = sh.text.strip()
            if not text: continue
            
            # Find questions in text
            matches = list(re.finditer(r'(?:^|\n)(\d+)[\.\s]+', text))
            if not matches: continue
            
            for m_i, m in enumerate(matches):
                start_pos = m.start()
                end_pos = matches[m_i+1].start() if m_i+1 < len(matches) else len(text)
                chunk = text[start_pos:end_pos].strip()
                
                m_q = re.match(r'^(\d+)[\.\s]+(.+)', chunk, re.DOTALL)
                if not m_q: continue
                q_num = int(m_q.group(1))
                body = m_q.group(2).strip()
                
                # split opts
                opts = []
                q_t = body
                if any(c in body for c in ['①', '②', '③', '④']):
                    parts = re.split(r'([①②③④])', body)
                    q_t = parts[0].strip()
                    for i in range(1, len(parts), 2):
                        ck = parts[i]
                        cv = parts[i+1].strip() if i+1 < len(parts) else ""
                        opts.append({"key": circle_map.get(ck, ck), "text": clean_text(cv)})
                else:
                    lines = [l.strip() for l in body.split('\n') if l.strip()]
                    if len(lines) > 1:
                        q_t = lines[0]
                        for idx, l in enumerate(lines[1:]):
                            m_opt = re.match(r'^([A-D1-4])[\.\s]+(.*)', l)
                            if m_opt:
                                k = m_opt.group(1)
                                if k in ['1','2','3','4']: k = chr(ord('A') + int(k) - 1)
                                opts.append({"key": k, "text": clean_text(m_opt.group(2))})
                            else:
                                opts.append({"key": chr(ord('A') + idx), "text": clean_text(l)})
                                
                ans_key = answers_map.get(q_num, "")
                
                exam_questions.append({
                    "id": f"{exam_id}_Q{q_num:02d}",
                    "exam_id": exam_id,
                    "section": sec_name,
                    "question_number": q_num,
                    "question": clean_text(q_t),
                    "options": opts,
                    "correct_answer": ans_key,
                    "images": slide_imgs if len(slide_imgs) > 0 else []
                })
                
    # Sort by question_number
    exam_questions.sort(key=lambda x: x['question_number'])
    
    exam_obj = {
        "id": exam_id,
        "title": exam_title,
        "total_questions": len(exam_questions),
        "questions": exam_questions
    }
    return exam_obj

# Extract De 01
e_de01 = parse_exam_pptx(
    os.path.join(BASE_DIR, "WS/04. De thi thu White Star/De thi thu White Star - De 01.pptx"),
    "EXAM_DE_01",
    "Đề thi thử White Star - Đề số 01",
    manual_answers=de01_answers
)
mock_exams.append(e_de01)

# Extract Test 01
e_test01 = parse_exam_pptx(
    os.path.join(BASE_DIR, "WS/04. De thi thu White Star/De thi thu WS Test 01.pptx"),
    "EXAM_TEST_01",
    "Đề thi thử White Star - Test 01"
)
mock_exams.append(e_test01)

# Extract Test 02
e_test02 = parse_exam_pptx(
    os.path.join(BASE_DIR, "WS/04. De thi thu White Star/De thi thu WS Test 02.pptx"),
    "EXAM_TEST_02",
    "Đề thi thử White Star - Test 02"
)
mock_exams.append(e_test02)

# Extract Test 03
e_test03 = parse_exam_pptx(
    os.path.join(BASE_DIR, "WS/04. De thi thu White Star/De thi thu WS Test 03.pptx"),
    "EXAM_TEST_03",
    "Đề thi thử White Star - Test 03"
)
mock_exams.append(e_test03)

# Extract Test 04
e_test04 = parse_exam_pptx(
    os.path.join(BASE_DIR, "WS/04. De thi thu White Star/De thi thu WS Test 04.pptx"),
    "EXAM_TEST_04",
    "Đề thi thử White Star - Test 04"
)
mock_exams.append(e_test04)

print(f"Parsed {len(mock_exams)} mock exams.")
for ex in mock_exams:
    print(f"  - {ex['title']}: {len(ex['questions'])} questions")

# Save files
with open(os.path.join(DATA_DIR, "questions_master.json"), "w", encoding="utf-8") as f:
    json.dump(master_questions, f, ensure_ascii=False, indent=2)

for k, v in by_subject.items():
    with open(os.path.join(SUBJECTS_DIR, f"{k}.json"), "w", encoding="utf-8") as f:
        json.dump(v, f, ensure_ascii=False, indent=2)

with open(os.path.join(DATA_DIR, "questions_review_bank.json"), "w", encoding="utf-8") as f:
    json.dump(review_questions, f, ensure_ascii=False, indent=2)

with open(os.path.join(DATA_DIR, "mock_exams.json"), "w", encoding="utf-8") as f:
    json.dump(mock_exams, f, ensure_ascii=False, indent=2)

for ex in mock_exams:
    ex_file = os.path.join(EXAMS_DIR, f"{ex['id'].lower()}.json")
    with open(ex_file, "w", encoding="utf-8") as f:
        json.dump(ex, f, ensure_ascii=False, indent=2)

# ==============================================================================
# 4. TẠO SUMMARY METADATA JSON
# ==============================================================================
# Count images
img_count = 0
for r, d, fs in os.walk(IMAGES_DIR):
    for f in fs:
        if any(f.endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.wmf']):
            img_count += 1

summary_data = {
    "project": "White Star Quiz Data Extraction",
    "generated_at": "2026-10-07",
    "total_master_questions": len(master_questions),
    "master_questions_by_subject": {
        "electric": len(by_subject['electric']),
        "plc": len(by_subject['plc']),
        "machine": len(by_subject['machine']),
        "pneumatics": len(by_subject['pneumatics'])
    },
    "difficulty_breakdown": {
        "Dễ": sum(1 for q in master_questions if q['difficulty'] == "Dễ"),
        "Trung bình": sum(1 for q in master_questions if q['difficulty'] == "Trung bình"),
        "Khó": sum(1 for q in master_questions if q['difficulty'] == "Khó")
    },
    "total_review_questions": len(review_questions),
    "total_mock_exams": len(mock_exams),
    "mock_exam_list": [
        {"id": ex['id'], "title": ex['title'], "total_questions": len(ex['questions'])}
        for ex in mock_exams
    ],
    "total_extracted_images": img_count,
    "directories": {
        "master_json": "data/questions_master.json",
        "by_subject": "data/questions_by_subject/",
        "review_bank_json": "data/questions_review_bank.json",
        "mock_exams": "data/exams/",
        "images": "data/images/"
    }
}

with open(os.path.join(DATA_DIR, "summary.json"), "w", encoding="utf-8") as f:
    json.dump(summary_data, f, ensure_ascii=False, indent=2)

print("\n=== ALL EXTRACTIONS COMPLETED SUCCESSFULLY! ===")
print(json.dumps(summary_data, ensure_ascii=False, indent=2))
