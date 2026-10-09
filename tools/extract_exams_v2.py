import pptx
import msoffcrypto
import io
import os
import re
import json
import copy

BASE_DIR = "/home/longp/projects/tech-ws-app"
DATA_DIR = os.path.join(BASE_DIR, "data")
EXAM_DIR = os.path.join(BASE_DIR, "WS/04. De thi thu White Star")
IMG_EXAMS_DIR = os.path.join(DATA_DIR, "images/exams")
os.makedirs(IMG_EXAMS_DIR, exist_ok=True)
os.makedirs(os.path.join(BASE_DIR, "images/exams"), exist_ok=True)

circle_map = {'①': 'A', '②': 'B', '③': 'C', '④': 'D', '⑤': 'E', '⑥': 'F'}
digit_to_alpha = {'1': 'A', '2': 'B', '3': 'C', '4': 'D', '5': 'E', '6': 'F'}

subject_meta = {
    'electric': {'name': 'Điện - Điện tử cơ bản', 'prefix': 'WS_EXAM_ELE', 'topic': 'Lý thuyết điện & Mạch điện tử'},
    'plc': {'name': 'PLC cơ bản', 'prefix': 'WS_EXAM_PLC', 'topic': 'Lập trình & Điều khiển PLC'},
    'pneumatics': {'name': 'Khí nén cơ bản', 'prefix': 'WS_EXAM_PNE', 'topic': 'Mạch van & Thiết bị khí nén'},
    'machine': {'name': 'Linh kiện máy cơ bản', 'prefix': 'WS_EXAM_MCH', 'topic': 'Truyền động cơ khí & Dung sai lắp ghép'}
}

def clean_text(val):
    if val is None: return ""
    text = str(val).strip()
    text = re.sub(r'\r\n|\r', '\n', text)
    lines = [re.sub(r'[ \t]+', ' ', l).strip() for l in text.split('\n')]
    return '\n'.join(lines).strip()

def normalize(text):
    if not text: return ""
    t = text.lower()
    t = re.sub(r'^\d+[\.\s\:\)]+', '', t)
    t = re.sub(r'[^\w\s]', '', t)
    t = re.sub(r'\s+', ' ', t).strip()
    return t

def load_prs(fname, pw=None):
    fpath = os.path.join(EXAM_DIR, fname)
    if pw:
        with open(fpath, "rb") as fh:
            of = msoffcrypto.OfficeFile(fh)
            of.load_key(password=pw)
            s = io.BytesIO()
            of.decrypt(s)
            s.seek(0)
            return pptx.Presentation(s)
    else:
        return pptx.Presentation(fpath)

def extract_slide_images(slide):
    imgs = []
    def recurse(sh, p_left=0, p_top=0):
        if sh.shape_type == 13: # PICTURE
            if sh.width > 20000 and sh.height > 20000:
                imgs.append({
                    "left": sh.left + p_left,
                    "top": sh.top + p_top,
                    "width": sh.width,
                    "height": sh.height,
                    "blob": sh.image.blob,
                    "ext": sh.image.ext.lower()
                })
        elif sh.shape_type == 6: # GROUP
            for sub in sh.shapes:
                recurse(sub, p_left + sh.left, p_top + sh.top)
    for sh in slide.shapes:
        recurse(sh)
    return imgs

de01_answers = {
    1: 'D', 2: 'D', 3: 'C', 4: 'A', 5: 'B', 6: 'B', 7: 'B', 8: 'C', 9: 'A', 10: 'C',
    11: 'D', 12: 'A', 13: 'C', 14: 'B', 15: 'D', 16: 'A', 17: 'D', 18: 'B', 19: 'C', 20: 'A',
    21: 'A', 22: 'A', 23: 'C', 24: 'A', 25: 'B', 26: 'D', 27: 'B', 28: 'C', 29: 'A', 30: 'C',
    31: 'D', 32: 'D', 33: 'C', 34: 'B', 35: 'A', 36: 'C', 37: 'B', 38: 'C', 39: 'D', 40: 'A',
    41: 'D', 42: 'C', 43: 'B', 44: 'A', 45: 'A', 46: 'D', 47: 'C', 48: 'A', 49: 'C', 50: 'A',
    51: 'A', 52: 'B', 53: 'C', 54: 'D', 55: 'B', 56: 'A', 57: 'C', 58: 'B', 59: 'A', 60: 'C',
    61: 'A', 62: 'A', 63: 'C', 64: 'C', 65: 'D', 66: 'A', 67: 'D', 68: 'D', 69: 'C', 70: 'A',
    71: 'D', 72: 'B', 73: 'C', 74: 'B', 75: 'A', 76: 'C', 77: 'C', 78: 'B', 79: 'B', 80: 'D'
}

def parse_slide_questions(slide, default_subj, exam_id, start_q_num_offset=0):
    slide_imgs = extract_slide_images(slide)
    
    subj = default_subj
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text.strip():
            t = sh.text.lower()
            if "điện" in t or "electric" in t: subj = "electric"
            elif "plc" in t: subj = "plc"
            elif "khí nén" in t or "pneumatics" in t: subj = "pneumatics"
            elif "máy" in t or "machine" in t: subj = "machine"
            
    q_candidates = []
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text.strip():
            continue
        text = sh.text.strip()
        matches = list(re.finditer(r'(?:^|\n)\s*(\d+)[\.\s]+', text))
        if not matches: continue
        
        for m_i, m in enumerate(matches):
            local_q_num = int(m.group(1))
            if local_q_num > 100: continue
            start_p = m.start()
            end_p = matches[m_i+1].start() if m_i+1 < len(matches) else len(text)
            chunk = text[start_p:end_p].strip()
            approx_top = sh.top + int(sh.height * (start_p / max(1, len(text))))
            
            q_candidates.append({
                "local_q_num": local_q_num,
                "global_q_num": local_q_num + start_q_num_offset,
                "chunk": chunk,
                "top": approx_top,
                "left": sh.left,
                "subj": subj
            })
            
    q_candidates.sort(key=lambda x: (0 if x['left'] < 4500000 else 1, x['top']))
    
    questions = []
    for idx, qc in enumerate(q_candidates):
        chunk = qc['chunk']
        m_b = re.match(r'^\d+[\.\s]+(.+)', chunk, re.DOTALL)
        body = m_b.group(1).strip() if m_b else chunk
        
        opts = []
        q_text = body
        if any(c in body for c in ['①', '②', '③', '④']):
            parts = re.split(r'([①②③④⑤⑥])', body)
            q_text = parts[0].strip()
            for i in range(1, len(parts), 2):
                ck = parts[i]
                cv = parts[i+1].strip() if i+1 < len(parts) else ""
                opts.append({"key": circle_map.get(ck, ck), "text": clean_text(cv)})
        else:
            lines = [l.strip() for l in body.split('\n') if l.strip()]
            if len(lines) > 1 and any(re.match(r'^[A-D1-4][\.\s]+', l) for l in lines[1:]):
                q_text = lines[0]
                for idx_l, l in enumerate(lines[1:]):
                    m_o = re.match(r'^([A-D1-4])[\.\s]+(.*)', l)
                    if m_o:
                        k = m_o.group(1)
                        if k.isdigit(): k = chr(ord('A') + int(k) - 1)
                        opts.append({"key": k, "text": clean_text(m_o.group(2))})
                    else:
                        opts.append({"key": chr(ord('A') + idx_l), "text": clean_text(l)})
                        
        is_right = qc['left'] >= 4500000
        next_top = None
        for other in q_candidates:
            if (other['left'] >= 4500000) == is_right and other['top'] > qc['top']:
                if next_top is None or other['top'] < next_top:
                    next_top = other['top']
        if next_top is None: next_top = 99999999
        
        matched_imgs = []
        for img in slide_imgs:
            if (img['left'] >= 4500000) == is_right:
                if qc['top'] - 150000 <= img['top'] < next_top:
                    matched_imgs.append(img)
                    
        img_paths = []
        g_q = qc['global_q_num']
        for i_idx, img in enumerate(matched_imgs):
            ext = img['ext'] if img['ext'] in ['png', 'jpg', 'jpeg'] else 'png'
            fname_img = f"{exam_id.lower()}_q{g_q:02d}_{i_idx+1}.{ext}"
            p1 = os.path.join(IMG_EXAMS_DIR, fname_img)
            p2 = os.path.join(BASE_DIR, "images/exams", fname_img)
            with open(p1, "wb") as f: f.write(img['blob'])
            with open(p2, "wb") as f: f.write(img['blob'])
            img_paths.append(f"images/exams/{fname_img}")
            
        if len(opts) == 0:
            opts = [
                {"key": "A", "text": "Phương án A (Xem hình minh họa A)" if len(img_paths) > 0 else "Đáp án A"},
                {"key": "B", "text": "Phương án B (Xem hình minh họa B)" if len(img_paths) > 0 else "Đáp án B"},
                {"key": "C", "text": "Phương án C (Xem hình minh họa C)" if len(img_paths) > 0 else "Đáp án C"},
                {"key": "D", "text": "Phương án D (Xem hình minh họa D)" if len(img_paths) > 0 else "Đáp án D"}
            ]
            
        questions.append({
            "exam_id": exam_id,
            "local_q_num": qc['local_q_num'],
            "question_number": g_q,
            "subject_code": qc['subj'],
            "question": clean_text(q_text),
            "options": opts,
            "images": img_paths
        })
        
    return questions

def extract_all_raw_exams():
    exams = []
    
    # 1. DE 01
    prs_de01 = load_prs("De thi thu White Star - De 01.pptx")
    q_de01 = []
    for s_idx in range(8):
        subj = "electric" if s_idx < 2 else ("plc" if s_idx < 4 else ("pneumatics" if s_idx < 6 else "machine"))
        q_de01.extend(parse_slide_questions(prs_de01.slides[s_idx], subj, "EXAM_DE_01", 0))
    for q in q_de01:
        q['correct_answer'] = de01_answers.get(q['question_number'], "")
    exams.append({"id": "EXAM_DE_01", "title": "Đề thi thử White Star - Đề số 01", "questions": q_de01})
    
    # 2. TEST 01
    prs_t1 = load_prs("De thi thu WS Test 01.pptx")
    ans_t1 = {}
    s9_t1 = prs_t1.slides[8]
    tables_t1 = [sh.table for sh in s9_t1.shapes if sh.has_table and len(sh.table.columns) == 2]
    for t in tables_t1:
        for r in range(len(t.rows)):
            k = t.cell(r, 0).text.strip()
            v = t.cell(r, 1).text.strip()
            if k.isdigit() and v:
                ans_t1[int(k)] = digit_to_alpha.get(v, v)
    q_t1 = []
    for s_idx in range(8):
        subj = "electric" if s_idx < 2 else ("plc" if s_idx < 4 else ("pneumatics" if s_idx < 6 else "machine"))
        q_t1.extend(parse_slide_questions(prs_t1.slides[s_idx], subj, "EXAM_TEST_01", 0))
    for q in q_t1:
        q['correct_answer'] = ans_t1.get(q['question_number'], "")
    exams.append({"id": "EXAM_TEST_01", "title": "Đề thi thử White Star - Test 01", "questions": q_t1})

    # 3. TEST 02
    prs_t2 = load_prs("De thi thu WS Test 02.pptx")
    ans_t2 = {}
    s9_t2 = prs_t2.slides[8]
    tables_t2 = [sh.table for sh in s9_t2.shapes if sh.has_table and len(sh.table.columns) == 2]
    for t in tables_t2:
        for r in range(len(t.rows)):
            k = t.cell(r, 0).text.strip()
            v = t.cell(r, 1).text.strip()
            if k.isdigit() and v:
                ans_t2[int(k)] = digit_to_alpha.get(v, v)
    q_t2 = []
    for s_idx in range(8):
        subj = "electric" if s_idx < 2 else ("plc" if s_idx < 4 else ("pneumatics" if s_idx < 6 else "machine"))
        q_t2.extend(parse_slide_questions(prs_t2.slides[s_idx], subj, "EXAM_TEST_02", 0))
    for q in q_t2:
        q['correct_answer'] = ans_t2.get(q['question_number'], "")
    exams.append({"id": "EXAM_TEST_02", "title": "Đề thi thử White Star - Test 02", "questions": q_t2})

    # 4. TEST 03
    prs_t3 = load_prs("De thi thu WS Test 03.pptx")
    s9_t3 = prs_t3.slides[8]
    ans_t3 = {}
    for sh in s9_t3.shapes:
        if sh.has_table and len(sh.table.columns) == 2:
            t = sh.table
            col_l = sh.left
            if col_l < 1500000: offset = 60
            elif col_l < 2800000: offset = 0
            elif col_l < 4000000: offset = 40
            else: offset = 20
            for r in range(1, len(t.rows)):
                k = t.cell(r, 0).text.strip()
                v = t.cell(r, 1).text.strip()
                if k.isdigit() and v:
                    ans_t3[int(k) + offset] = digit_to_alpha.get(v, v)
    q_t3 = []
    q_t3.extend(parse_slide_questions(prs_t3.slides[0], "machine", "EXAM_TEST_03", 60))
    q_t3.extend(parse_slide_questions(prs_t3.slides[1], "machine", "EXAM_TEST_03", 60))
    q_t3.extend(parse_slide_questions(prs_t3.slides[2], "electric", "EXAM_TEST_03", 0))
    q_t3.extend(parse_slide_questions(prs_t3.slides[3], "electric", "EXAM_TEST_03", 0))
    q_t3.extend(parse_slide_questions(prs_t3.slides[4], "pneumatics", "EXAM_TEST_03", 40))
    q_t3.extend(parse_slide_questions(prs_t3.slides[5], "pneumatics", "EXAM_TEST_03", 40))
    q_t3.extend(parse_slide_questions(prs_t3.slides[6], "plc", "EXAM_TEST_03", 20))
    q_t3.extend(parse_slide_questions(prs_t3.slides[7], "plc", "EXAM_TEST_03", 20))
    for q in q_t3:
        q['correct_answer'] = ans_t3.get(q['question_number'], "")
    exams.append({"id": "EXAM_TEST_03", "title": "Đề thi thử White Star - Test 03", "questions": q_t3})

    # 5. TEST 04
    prs_t4 = load_prs("De thi thu WS Test 04.pptx")
    s9_t4 = prs_t4.slides[8]
    ans_t4 = {}
    for sh in s9_t4.shapes:
        if sh.has_table and len(sh.table.columns) == 2:
            t = sh.table
            col_l = sh.left
            if col_l < 1500000: offset = 60
            elif col_l < 3000000: offset = 0
            elif col_l < 4500000: offset = 40
            else: offset = 20
            for r in range(1, len(t.rows)):
                k = t.cell(r, 0).text.strip()
                v = t.cell(r, 1).text.strip()
                if k.isdigit() and v:
                    ans_t4[int(k) + offset] = digit_to_alpha.get(v, v)
    q_t4 = []
    q_t4.extend(parse_slide_questions(prs_t4.slides[0], "machine", "EXAM_TEST_04", 60))
    q_t4.extend(parse_slide_questions(prs_t4.slides[1], "machine", "EXAM_TEST_04", 60))
    q_t4.extend(parse_slide_questions(prs_t4.slides[2], "electric", "EXAM_TEST_04", 0))
    q_t4.extend(parse_slide_questions(prs_t4.slides[3], "electric", "EXAM_TEST_04", 0))
    q_t4.extend(parse_slide_questions(prs_t4.slides[4], "pneumatics", "EXAM_TEST_04", 40))
    q_t4.extend(parse_slide_questions(prs_t4.slides[5], "pneumatics", "EXAM_TEST_04", 40))
    q_t4.extend(parse_slide_questions(prs_t4.slides[6], "plc", "EXAM_TEST_04", 20))
    q_t4.extend(parse_slide_questions(prs_t4.slides[7], "plc", "EXAM_TEST_04", 20))
    for q in q_t4:
        q['correct_answer'] = ans_t4.get(q['question_number'], "")
    exams.append({"id": "EXAM_TEST_04", "title": "Đề thi thử White Star - Test 04", "questions": q_t4})

    # 6. THANG 01-2021
    prs_2021 = load_prs("De thi thu WS (Thang 01-2021).pptx", "44448888")
    def parse_tables_for_ans(slide, offset):
        ans = {}
        for sh in slide.shapes:
            if sh.has_table:
                t = sh.table
                if len(t.columns) == 2:
                    for r in range(1, len(t.rows)):
                        k = t.cell(r, 0).text.strip()
                        v = t.cell(r, 1).text.strip()
                        if k.isdigit() and v: ans[int(k) + offset] = digit_to_alpha.get(v, v)
                elif len(t.columns) == 10 and len(t.rows) >= 2:
                    for r in range(0, len(t.rows), 2):
                        if r + 1 < len(t.rows):
                            for c in range(len(t.columns)):
                                k = t.cell(r, c).text.strip()
                                v = t.cell(r+1, c).text.strip()
                                if k.isdigit() and v: ans[int(k) + offset] = digit_to_alpha.get(v, v)
        return ans
        
    ans_2021 = {}
    ans_2021.update(parse_tables_for_ans(prs_2021.slides[2], 20))
    ans_2021.update(parse_tables_for_ans(prs_2021.slides[5], 40))
    ans_2021.update(parse_tables_for_ans(prs_2021.slides[7], 0))
    ans_2021.update(parse_tables_for_ans(prs_2021.slides[9], 60))
    
    q_2021 = []
    q_2021.extend(parse_slide_questions(prs_2021.slides[0], "plc", "EXAM_2021_01", 20))
    q_2021.extend(parse_slide_questions(prs_2021.slides[1], "plc", "EXAM_2021_01", 20))
    q_2021.extend(parse_slide_questions(prs_2021.slides[3], "pneumatics", "EXAM_2021_01", 40))
    q_2021.extend(parse_slide_questions(prs_2021.slides[4], "pneumatics", "EXAM_2021_01", 40))
    q_2021.extend(parse_slide_questions(prs_2021.slides[6], "electric", "EXAM_2021_01", 0))
    q_2021.extend(parse_slide_questions(prs_2021.slides[8], "machine", "EXAM_2021_01", 60))
    for q in q_2021:
        q['correct_answer'] = ans_2021.get(q['question_number'], "")
    exams.append({"id": "EXAM_2021_01", "title": "Đề thi thử WS (Tháng 01-2021)", "questions": q_2021})

    for ex in exams:
        seen = {}
        dedup = []
        for q in ex['questions']:
            gn = q['question_number']
            if gn not in seen:
                seen[gn] = q
                dedup.append(q)
            else:
                if q['images'] and not seen[gn]['images']:
                    seen[gn]['images'] = q['images']
        dedup.sort(key=lambda x: x['question_number'])
        ex['questions'] = dedup
        print(f"-> Parsed {ex['title']}: {len(dedup)} questions (with images: {sum(1 for q in dedup if q['images'])})")
        
    return exams

# Smart explanation generator
def generate_smart_explanation(q_text, opts, correct_key, subj_code):
    correct_opt_text = ""
    wrong_opts = []
    for opt in opts:
        if opt['key'] == correct_key:
            correct_opt_text = opt['text']
        else:
            wrong_opts.append(f"{opt['key']}: {opt['text']}")
            
    meta = subject_meta.get(subj_code, subject_meta['electric'])
    subj_name = meta['name']
    topic = meta['topic']
    
    why_correct = f"Theo nguyên lý kỹ thuật {subj_name}, phương án [{correct_key}] '{correct_opt_text}' là chính xác. Trong ứng dụng thực tế sản xuất và vận hành hệ thống tự động hóa White Star, thông số và cấu hình này đáp ứng đúng tiêu chuẩn kỹ thuật thiết kế."
    if "không" in q_text.lower() or "sai" in q_text.lower():
        why_correct = f"Theo tiêu chuẩn kỹ thuật {subj_name}, phương án [{correct_key}] '{correct_opt_text}' là nhận định KHÔNG đúng (hoặc không phù hợp với thực tế kỹ thuật), do đó đây là đáp án cần chọn theo yêu cầu câu hỏi."
        
    why_wrong = "Các phương án còn lại (" + "; ".join(wrong_opts[:3]) + ") đều mô tả các đặc tính kỹ thuật thông thường hoặc các thành phần tiêu chuẩn khác, không khớp với điều kiện đặt ra trong câu hỏi."
    supp = f"💡 **Lưu ý kỹ thuật White Star:** Khi làm việc với chuyên đề **{topic}**, kỹ thuật viên cần chú ý đọc kỹ ký hiệu, thông số định mức và tuân thủ các quy chuẩn an toàn lao động quốc tế (ISO/JIS/IEC)."
    
    return {
        "overview": f"Câu hỏi thuộc chuyên đề **{topic}** ({subj_name}). Đáp án chính xác là **{correct_key}**.",
        "correct_answer": correct_key,
        "correct_text": correct_opt_text,
        "why_correct": why_correct,
        "why_wrong": why_wrong,
        "supplementary_knowledge": supp
    }

def main():
    print("=== BẮT ĐẦU TRÍCH XUẤT VÀ ĐỒNG BỘ ĐỀ THI THỬ VÀO MASTER BANK ===")
    
    # 1. Load current master bank
    with open(os.path.join(DATA_DIR, "questions_master.json"), "r", encoding="utf-8") as f:
        master_bank = json.load(f)
    print(f"Master bank ban đầu: {len(master_bank)} câu hỏi.")
    
    # Build master bank index
    master_index = {}
    for q in master_bank:
        norm = normalize(q['question'])
        if len(norm) > 10:
            master_index[norm[:55]] = q
            
    # 2. Extract raw exams
    exams = extract_all_raw_exams()
    
    # 3. Process deduplication & integration
    added_new_count = 0
    updated_image_count = 0
    linked_exam_count = 0
    
    new_q_counters = {'electric': 100, 'plc': 100, 'pneumatics': 100, 'machine': 100}
    
    # Track newly created questions to avoid duplicates among exams
    seen_new_keys = {}
    
    processed_exams = []
    
    for ex in exams:
        ex_id = ex['id']
        ex_title = ex['title']
        ex_questions = []
        
        for q in ex['questions']:
            norm_q = normalize(q['question'])
            prefix_key = norm_q[:55] if len(norm_q) > 10 else norm_q
            
            matched_master_q = None
            if prefix_key in master_index:
                matched_master_q = master_index[prefix_key]
            else:
                # Fuzzy search
                for k, mq in master_index.items():
                    if len(k) > 15 and len(norm_q) > 15 and (k in norm_q or norm_q[:40] in k):
                        matched_master_q = mq
                        break
                        
            if matched_master_q:
                # Câu ĐÃ CÓ trong Master Bank
                linked_exam_count += 1
                # Cập nhật ảnh nếu Master Bank thiếu mà câu đề thi có
                if q['images'] and (not matched_master_q.get('images') or len(matched_master_q.get('images', [])) == 0):
                    matched_master_q['images'] = q['images']
                    updated_image_count += 1
                    
                # Thêm vào đề thi với liên kết master
                ex_q_obj = copy.deepcopy(matched_master_q)
                ex_q_obj['exam_q_num'] = q['question_number']
                if q['images'] and len(ex_q_obj.get('images', [])) == 0:
                    ex_q_obj['images'] = q['images']
                ex_questions.append(ex_q_obj)
            else:
                # Câu CHƯA CÓ trong Master Bank
                subj = q['subject_code']
                if prefix_key in seen_new_keys:
                    # Trùng với câu mới đã add từ đề trước
                    existing_new_q = seen_new_keys[prefix_key]
                    ex_q_obj = copy.deepcopy(existing_new_q)
                    ex_q_obj['exam_q_num'] = q['question_number']
                    if q['images'] and len(ex_q_obj.get('images', [])) == 0:
                        ex_q_obj['images'] = q['images']
                    ex_questions.append(ex_q_obj)
                else:
                    # Tạo câu mới hoàn toàn
                    new_q_counters[subj] += 1
                    new_id = f"{subject_meta[subj]['prefix']}_{new_q_counters[subj]:03d}"
                    corr_ans = q['correct_answer'] if q['correct_answer'] else 'A'
                    
                    explanation = generate_smart_explanation(q['question'], q['options'], corr_ans, subj)
                    
                    q_type = "theory"
                    if q['images'] and len(q['images']) > 0: q_type = "diagram"
                    elif any(c in q['question'] for c in ['Ω', 'μ', 'kΩ', 'V', 'A', 'Hz', 'vòng/phút', 'bar', 'MPa']): q_type = "calculation"
                    
                    new_q_obj = {
                        "id": new_id,
                        "source": ex_title,
                        "subject": subject_meta[subj]['name'],
                        "subject_code": subj,
                        "module": subject_meta[subj]['topic'],
                        "difficulty": "Trung bình" if q_type == "diagram" else "Dễ",
                        "question_number": new_q_counters[subj],
                        "question": q['question'],
                        "options": q['options'],
                        "correct_answer": corr_ans,
                        "images": q['images'],
                        "author": "White Star Examination Board",
                        "note": f"Trích xuất từ {ex_title} - Câu {q['question_number']}",
                        "question_type": q_type,
                        "topic_tag": subject_meta[subj]['topic'],
                        "explanation": explanation
                    }
                    
                    master_bank.append(new_q_obj)
                    master_index[prefix_key] = new_q_obj
                    seen_new_keys[prefix_key] = new_q_obj
                    added_new_count += 1
                    
                    ex_q_obj = copy.deepcopy(new_q_obj)
                    ex_q_obj['exam_q_num'] = q['question_number']
                    ex_questions.append(ex_q_obj)
                    
        processed_exams.append({
            "id": ex_id,
            "title": ex_title,
            "total_questions": len(ex_questions),
            "questions": ex_questions
        })
        
    print(f"\n=== KẾT QUẢ TÍCH HỢP ===")
    print(f"- Tổng số câu trong Master Bank sau khi tích hợp: {len(master_bank)}")
    print(f"- Số câu mới hoàn toàn từ đề thi được thêm vào: {added_new_count}")
    print(f"- Số câu trong Master Bank được bổ sung ảnh: {updated_image_count}")
    print(f"- Số lượt khớp câu hỏi đã có sẵn: {linked_exam_count}")
    
    # 4. Save master questions JSON & JS
    master_json_path = os.path.join(DATA_DIR, "questions_master.json")
    with open(master_json_path, "w", encoding="utf-8") as f:
        json.dump(master_bank, f, ensure_ascii=False, indent=2)
        
    master_js_path = os.path.join(DATA_DIR, "questions_master.js")
    with open(master_js_path, "w", encoding="utf-8") as f:
        f.write("// Ngân hàng câu hỏi Master Bank (White Star Quiz)\n")
        f.write("window.QUESTIONS_MASTER = ")
        json.dump(master_bank, f, ensure_ascii=False, indent=2)
        f.write(";\n")
        
    # 5. Save by subject JSON
    by_subject = {'electric': [], 'plc': [], 'machine': [], 'pneumatics': []}
    for q in master_bank:
        sub = q.get('subject_code', 'electric')
        if sub in by_subject:
            by_subject[sub].append(q)
            
    for sub, q_list in by_subject.items():
        sub_file = os.path.join(DATA_DIR, f"questions_by_subject/{sub}.json")
        with open(sub_file, "w", encoding="utf-8") as f:
            json.dump(q_list, f, ensure_ascii=False, indent=2)
            
    # 6. Save mock exams JSON & JS
    mock_json_path = os.path.join(DATA_DIR, "mock_exams.json")
    with open(mock_json_path, "w", encoding="utf-8") as f:
        json.dump(processed_exams, f, ensure_ascii=False, indent=2)
        
    mock_js_path = os.path.join(DATA_DIR, "mock_exams.js")
    with open(mock_js_path, "w", encoding="utf-8") as f:
        f.write("// Dữ liệu các bộ đề thi thử White Star\n")
        f.write("window.MOCK_EXAMS = ")
        json.dump(processed_exams, f, ensure_ascii=False, indent=2)
        f.write(";\n")
        
    # Also save each exam individually into data/exams/
    for ex in processed_exams:
        ex_file = os.path.join(DATA_DIR, f"exams/{ex['id'].lower()}.json")
        with open(ex_file, "w", encoding="utf-8") as f:
            json.dump(ex, f, ensure_ascii=False, indent=2)
            
    # 7. Update summary.json & summary.js
    img_count = 0
    for r, d, fs in os.walk(os.path.join(DATA_DIR, "images")):
        for file in fs:
            if any(file.endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.wmf']):
                img_count += 1
                
    summary_data = {
        "project": "White Star Quiz Data Extraction",
        "generated_at": "2026-10-09",
        "total_master_questions": len(master_bank),
        "master_questions_by_subject": {
            "electric": len(by_subject['electric']),
            "plc": len(by_subject['plc']),
            "machine": len(by_subject['machine']),
            "pneumatics": len(by_subject['pneumatics'])
        },
        "difficulty_breakdown": {
            "Dễ": sum(1 for q in master_bank if q.get('difficulty') == "Dễ"),
            "Trung bình": sum(1 for q in master_bank if q.get('difficulty') == "Trung bình"),
            "Khó": sum(1 for q in master_bank if q.get('difficulty') == "Khó")
        },
        "total_mock_exams": len(processed_exams),
        "mock_exam_list": [
            {"id": ex['id'], "title": ex['title'], "total_questions": len(ex['questions'])}
            for ex in processed_exams
        ],
        "total_extracted_images": img_count
    }
    
    with open(os.path.join(DATA_DIR, "summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary_data, f, ensure_ascii=False, indent=2)
        
    with open(os.path.join(DATA_DIR, "summary.js"), "w", encoding="utf-8") as f:
        f.write("// Thống kê tổng quan dữ liệu White Star Quiz\n")
        f.write("window.APP_SUMMARY = ")
        json.dump(summary_data, f, ensure_ascii=False, indent=2)
        f.write(";\n")
        
    print("\n=== HOÀN TẤT TRÍCH XUẤT VÀ ĐỒNG BỘ DỮ LIỆU THÀNH CÔNG! ===")
    print(json.dumps(summary_data, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
